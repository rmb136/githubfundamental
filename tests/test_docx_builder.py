import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt

from github_im.content.catalog import build_module
from github_im.docx_builder import build_docx


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "figures" / "cmu-module"


class DocxBuilderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.path = Path(cls.temp_dir.name) / "module.docx"
        build_docx(build_module(), ASSETS, cls.path)
        cls.document = Document(cls.path)
        cls.text = "\n".join(p.text for p in cls.document.paragraphs)

    @classmethod
    def tearDownClass(cls):
        cls.temp_dir.cleanup()

    def test_uses_cmu_page_geometry_and_body_style(self):
        section = self.document.sections[0]
        self.assertAlmostEqual(section.page_width, Mm(210), delta=1000)
        self.assertAlmostEqual(section.page_height, Mm(297), delta=1000)
        self.assertEqual(section.top_margin, Inches(1))
        self.assertEqual(section.bottom_margin, Inches(1))
        self.assertEqual(section.left_margin, Inches(1.5))
        self.assertEqual(section.right_margin, Inches(1))
        normal = self.document.styles["Normal"]
        self.assertEqual(normal.font.name, "Arial")
        self.assertEqual(normal.font.size, Pt(11))
        self.assertEqual(normal.paragraph_format.line_spacing, 1)
        self.assertTrue(self.document.styles["Command"].paragraph_format.keep_together)

    def test_contains_all_required_sections(self):
        required = [
            "Preface", "Table of Contents", "List of Figures", "Glossary",
            "Semester Roadmap", "Capstone Project", "Assessment Rubrics",
            "Answer Key", "Git Command Cheat Sheet", "References",
            "About the Author",
        ]
        required.extend(f"Unit {number}:" for number in range(1, 14))
        for heading in required:
            with self.subTest(heading=heading):
                self.assertIn(heading, self.text)

    def test_cover_uses_requested_title_and_author(self):
        title_paragraph = next(
            paragraph for paragraph in self.document.paragraphs
            if paragraph.style.name == "Title"
        )
        self.assertEqual(title_paragraph.text, "GitHub Fundamentals for Beginners")
        self.assertIn("by RINANTE M. BUNTOD", self.text)

    def test_cover_title_is_black_and_has_no_decorative_border(self):
        title_paragraph = next(
            paragraph for paragraph in self.document.paragraphs
            if paragraph.style.name == "Title"
        )
        self.assertEqual(title_paragraph.runs[0].font.color.rgb, None)
        style_properties = self.document.styles["Title"].element.pPr
        self.assertIsNone(style_properties.find(qn("w:pBdr")))

    def test_cover_hides_running_furniture_and_includes_accessible_artwork(self):
        section = self.document.sections[0]
        self.assertTrue(section.different_first_page_header_footer)
        self.assertEqual(section.first_page_header.paragraphs[0].text, "")
        self.assertEqual(section.first_page_footer.paragraphs[0].text, "")
        with zipfile.ZipFile(self.path) as package:
            xml = package.read("word/document.xml").decode("utf-8")
        self.assertIn('descr="Abstract Git learning journey', xml)

    def test_author_identity_is_consistent_in_metadata_and_back_matter(self):
        self.assertEqual(self.document.core_properties.title, "GitHub Fundamentals for Beginners")
        self.assertEqual(self.document.core_properties.author, "Rinante M. Buntod")
        self.assertIn("Rinante M. Buntod is the author of this instructional module.", self.text)

    def test_major_content_headings_start_on_new_pages_without_break_paragraphs(self):
        targets = {f"Unit {number}:" for number in range(1, 14)} | {
            "Capstone Project", "Answer Key", "About the Author"
        }
        matches = [
            paragraph for paragraph in self.document.paragraphs
            if paragraph.style.name == "Heading 1"
            and any(paragraph.text.startswith(target) for target in targets)
        ]
        self.assertGreaterEqual(len(matches), 16)
        self.assertTrue(all(paragraph.paragraph_format.page_break_before for paragraph in matches))

    def test_embeds_thirteen_figures_with_captions_and_alt_text(self):
        captions = [
            p.text for p in self.document.paragraphs
            if p.style.name == "Caption" and p.text.startswith("Figure ")
        ]
        self.assertEqual(len(captions), 13)
        for number in range(1, 14):
            self.assertTrue(any(text.startswith(f"Figure {number}.") for text in captions))
        with zipfile.ZipFile(self.path) as package:
            xml = package.read("word/document.xml").decode("utf-8")
        self.assertEqual(xml.count("descr="), 14)
        self.assertNotIn('descr=""', xml)

    def test_tables_prevent_rows_from_splitting_across_pages(self):
        with zipfile.ZipFile(self.path) as package:
            xml = package.read("word/document.xml").decode("utf-8")
        self.assertGreaterEqual(xml.count("w:cantSplit"), 50)

    def test_each_instructional_numbered_list_restarts_at_one(self):
        with zipfile.ZipFile(self.path) as package:
            numbering = package.read("word/numbering.xml").decode("utf-8")
        self.assertGreaterEqual(numbering.count("w:startOverride"), 60)


if __name__ == "__main__":
    unittest.main()
