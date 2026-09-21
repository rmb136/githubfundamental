#!/usr/bin/env python3
"""
Generate an enhanced Word (.docx) and PDF for "GitHub Fundamentals: A Complete Instructional Guide for Students".

Features:
- Improved chapter explanations
- Sample illustrative images per chapter (generated with Pillow)
- Labeled figure captions
- Post-assessment 10-question quiz per chapter
- Attempts PDF conversion using docx2pdf (requires Word on Windows)

Run:
  python generate_github_fundamentals_docx.py

Dependencies: python-docx, pillow, docx2pdf (optional).
"""

import sys
import subprocess
import os


def ensure_package(pkg):
    try:
        __import__(pkg)
    except Exception:
        print(f"Installing missing package: {pkg}")
        subprocess.check_call([sys.executable, "-m", "pip", "install", pkg])


ensure_package("docx")
ensure_package("Pillow")
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


def add_heading(document, text, level=1):
    document.add_heading(text, level=level)


def add_paragraph(document, text, bold=False, style=None):
    p = document.add_paragraph(style=style) if style else document.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def add_page_header_footer(doc, header_text=None):
    section = doc.sections[0]
    if header_text:
        header = section.header
        header_para = header.paragraphs[0]
        header_para.text = header_text
        header_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = footer_para.add_run()
    # Insert page number field
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


def ensure_pillow():
    try:
        from PIL import Image  # noqa: F401
    except Exception:
        print("Pillow not found. Installing...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "pillow"])


def create_sample_image(path, title, label):
    ensure_pillow()
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1200, 675
    bg = Image.new("RGB", (width, height), color=(245, 245, 250))
    draw = ImageDraw.Draw(bg)

    # Header bar
    draw.rectangle([(0, 0), (width, 80)], fill=(30, 90, 150))
    try:
        font_h = ImageFont.truetype("arial.ttf", 32)
        font = ImageFont.truetype("arial.ttf", 24)
        small = ImageFont.truetype("arial.ttf", 16)
    except Exception:
        font_h = ImageFont.load_default()
        font = ImageFont.load_default()
        small = ImageFont.load_default()

    draw.text((20, 20), title, fill="white", font=font_h)

    # Simple diagram: local -> arrow -> remote
    draw.rectangle([(100, 150), (420, 300)], fill="white", outline="black")
    draw.text((140, 190), "Local Repository", fill="black", font=font)
    draw.polygon([(440, 205), (520, 240), (440, 275)], fill="black")
    draw.rectangle([(540, 150), (860, 300)], fill="white", outline="black")
    draw.text((580, 190), "Remote (GitHub)", fill="black", font=font)

    # Footer label
    draw.text((20, height - 30), f"Figure: {label}", fill=(80, 80, 80), font=small)

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    bg.save(path)


def chapter_quiz_questions(ch_num):
    q = []
    if ch_num == 1:
        q = [
            "What is Git and why is it used?",
            "How does Git differ from GitHub?",
            "Define a repository in your own words.",
            "What information does a commit store?",
            "Explain what a branch is.",
            "What does cloning a repository do?",
            "Name two benefits of version control.",
            "What is the purpose of a commit message?",
            "How would you view the commit history of a repo?",
            "Give an example scenario to create a branch.",
        ]
    elif ch_num == 2:
        q = [
            "How do you install Git on a machine? (list common methods)",
            "What commands set your Git username and email?",
            "Why is configuring your email important for commits?",
            "What is GitHub Desktop used for?",
            "Describe one benefit of the GitHub CLI.",
            "How do you connect your local Git to GitHub?",
            "What are the pros/cons of SSH vs HTTPS for Git remotes?",
            "How can you verify your Git installation?",
            "Where are global Git configs stored?",
            "How do you create a GitHub account?",
        ]
    else:
        for i in range(1, 11):
            q.append(
                f"Chapter {ch_num} question {i}: Explain or describe a concept from chapter {ch_num}."
            )
    return q


def add_chapter(doc, num, title, explanation_paragraphs, image_label):
    # Add a chapter with explanation, sample image, caption, and post-assessment quiz
    doc.add_page_break()
    add_heading(doc, f"Chapter {num}: {title}", level=1)
    for para in explanation_paragraphs:
        add_paragraph(doc, para)

    # Create and insert sample image
    images_dir = os.path.join(os.getcwd(), "figures")
    os.makedirs(images_dir, exist_ok=True)
    img_path = os.path.join(images_dir, f"fig_ch{num}.png")
    create_sample_image(img_path, f"Figure for Chapter {num}", image_label)
    try:
        doc.add_picture(img_path, width=Inches(6))
        cap = doc.add_paragraph(f"Figure {num}: {image_label}")
        cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    except Exception:
        add_paragraph(doc, f"[Image: {image_label}]")

    # Post-assessment quiz
    add_heading(doc, "Post-Assessment Quiz", level=2)
    questions = chapter_quiz_questions(num)
    for i, q in enumerate(questions, start=1):
        add_paragraph(doc, f"{i}) {q}")


def main():
    doc = Document()

    doc.add_heading(
        "GitHub Fundamentals: A Complete Instructional Guide for Students", 0
    )

    # Table of Contents (as numbered list)
    doc.add_heading("Table of Contents", level=1)
    toc_items = [
        "Introduction",
        "Chapter 1: Introduction to Git and GitHub",
        "Chapter 2: Setting Up Git and GitHub",
        "Chapter 3: Your First GitHub Repository",
        "Chapter 4: The Essential Git Commands",
        "Chapter 5: Branching and Merging",
        "Chapter 6: Collaborating with Teams",
        "Chapter 7: Managing Issues and Projects",
        "Chapter 8: GitHub Best Practices",
        "Chapter 9: GitHub Actions & Automation",
        "Chapter 10: GitHub Pages and Hosting",
        "Chapter 11: Advanced GitHub Features",
        "Chapter 12: Contributing to Open Source Projects",
        "Chapter 13: Troubleshooting and Common Issues",
        "Conclusion",
        "Appendix: Glossary, Resources, Cheat Sheet",
        "Teaching Strategy",
    ]
    for item in toc_items:
        para = doc.add_paragraph(item)
        try:
            para.style = "List Number"
        except Exception:
            pass

    # Introduction
    doc.add_page_break()
    add_heading(doc, "Introduction", level=1)
    add_paragraph(doc, "Why Learn GitHub?")
    add_paragraph(
        doc,
        "GitHub is the industry standard for hosting code, managing version control, and collaborating across teams. Students who learn GitHub gain practical skills for real-world development workflows, open-source contribution, and reproducible project management.",
    )

    # Chapters with improved explanations, figures, and quizzes
    chapters = [
        (
            1,
            "Introduction to Git and GitHub",
            [
                "Git is a distributed version control system that tracks changes to files over time, allowing you to revert to earlier versions, collaborate safely, and maintain a history of work.",
                "GitHub is a cloud-based hosting service built on top of Git. It adds collaboration features such as pull requests, issues, and web hosting.",
                "Key terms: repository (a project storage), commit (a recorded change), branch (parallel line of development), clone (a local copy of a repo).",
            ],
            "Local-vs-Remote-Overview",
        ),
        (
            2,
            "Setting Up Git and GitHub",
            [
                "Install Git from the official website, package manager, or use GitHub Desktop for a GUI experience.",
                "Configure your username and email so your commits are associated with your identity: these settings appear in commit metadata and on GitHub profiles.",
                "Set up SSH keys or use HTTPS for authentication; SSH is more convenient for frequent pushes.",
            ],
            "Git-Setup-Example",
        ),
        (
            3,
            "Your First GitHub Repository",
            [
                "A repository contains the project files, history, and configuration. On GitHub you can create repos from the website or import existing code.",
                "Cloning copies a remote repository to your local machine. Initializing a repo locally creates a new Git history for your project.",
                "Making your first commit captures the initial state of your project and starts the history.",
            ],
            "Create-Clone-Commit",
        ),
        (
            4,
            "The Essential Git Commands",
            [
                "git init initializes a new local repository.",
                "git add stages changes to be included in the next commit; git commit records those changes.",
                "git status shows current state; git log lists commit history; git push uploads commits to a remote; git pull fetches and merges updates.",
            ],
            "Core-Git-Commands",
        ),
        (
            5,
            "Branching and Merging",
            [
                "Branches let you work on features independently of the main codebase. Use branches for new features, bug fixes, or experiments.",
                "Merging integrates changes from one branch into another. Conflicts may occur when the same lines were edited; resolving them is part of collaboration.",
            ],
            "Branching-Merging",
        ),
        (
            6,
            "Collaborating with Teams",
            [
                "GitHub provides collaboration tools: forks, pull requests, code review, protected branches, and collaborators/teams.",
                "A typical workflow: fork (or branch) -> implement changes -> open a PR -> review -> merge when approved.",
            ],
            "Fork-PR-Workflow",
        ),
        (
            7,
            "Managing Issues and Projects",
            [
                "Issues track bugs, tasks, and feature requests. Labels, assignees, and milestones help triage work.",
                "GitHub Projects (Kanban) help visualize workflow: backlog -> in progress -> done.",
            ],
            "Issues-Project-Board",
        ),
        (
            8,
            "GitHub Best Practices",
            [
                "Write clear commit messages (what and why), keep repositories modular, maintain a concise README, and use .gitignore to exclude generated files.",
                "Use branches, descriptive PR titles, and periodic rebase or squash to keep history readable where appropriate.",
            ],
            "Best-Practices",
        ),
        (
            9,
            "GitHub Actions & Automation",
            [
                "Actions automate workflows like running tests, linting, and deployment. Define workflows in YAML files under .github/workflows.",
                "Continuous Integration (CI) runs tests automatically on push or PRs to catch regressions early.",
            ],
            "Actions-CI-Example",
        ),
        (
            10,
            "GitHub Pages and Hosting",
            [
                "GitHub Pages hosts static websites directly from a repository. Use it for project pages, documentation, or personal sites.",
                "Choose a publishing source (branch or /docs folder) and optionally configure a custom domain.",
            ],
            "GitHub-Pages-Example",
        ),
        (
            11,
            "Advanced GitHub Features",
            [
                "GitHub Packages lets you publish packages; Insights provides repository analytics; the GitHub API enables automation and integrations.",
                "Explore Actions marketplace, Dependabot alerts, and security scanning for mature workflows.",
            ],
            "Advanced-Features",
        ),
        (
            12,
            "Contributing to Open Source Projects",
            [
                "Find projects by interests, check contribution guidelines, and start with small issues or documentation improvements.",
                "Respect project etiquette: read the README and CODE_OF_CONDUCT, write clear PR descriptions, and respond to review feedback.",
            ],
            "Open-Source-Contrib",
        ),
        (
            13,
            "Troubleshooting and Common Issues",
            [
                "Common problems: merge conflicts, authentication errors, accidental commits. Use git reflog, revert, and reset carefully to recover.",
                "When push fails, check remote URL, authentication, and if the branch is protected.",
            ],
            "Troubleshooting",
        ),
    ]

    for num, title, paras, fig_label in chapters:
        add_chapter(doc, num, title, paras, fig_label)

    # Appendix
    doc.add_page_break()
    add_heading(doc, "Appendix", level=1)
    add_paragraph(doc, "Glossary of Terms")
    add_paragraph(doc, "Additional Resources (Books, Websites, Tutorials)")
    add_paragraph(doc, "Further Reading on GitHub Features")
    add_paragraph(doc, "Cheat Sheet of Common Git Commands")

    # Teaching Strategy
    doc.add_page_break()
    add_heading(doc, "Teaching Strategy", level=1)
    add_paragraph(
        doc,
        "Hands-On Approach: Each chapter will come with practical exercises and a post-assessment quiz.",
    )
    add_paragraph(
        doc,
        "Project-Based Learning: Students will work on collaborative projects and apply what they've learned.",
    )
    add_paragraph(
        doc,
        "Real-Life Scenarios: Encourage students to create personal or group projects on GitHub.",
    )
    add_paragraph(
        doc,
        "Peer Review & Collaboration: Foster a collaborative environment by integrating peer reviews and GitHub PRs.",
    )

    # Add header/footer and save
    add_page_header_footer(
        doc, header_text="GitHub Fundamentals: A Complete Instructional Guide"
    )
    from datetime import datetime

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_name = f"GitHub_Fundamentals_{ts}.docx"
    doc.save(out_name)
    print(f"Created {out_name}")

    # Attempt PDF conversion using docx2pdf (Word must be installed)
    try:
        ensure_package("docx2pdf")
        from docx2pdf import convert

        pdf_name = f"GitHub_Fundamentals_{ts}.pdf"
        print("Converting to PDF...")
        convert(out_name, pdf_name)
        print(f"Created {pdf_name}")
    except Exception as e:
        print("PDF conversion skipped or failed:", e)
        print(
            "If you need a PDF, ensure Microsoft Word is installed on Windows or use LibreOffice to convert the .docx file."
        )

    # Create an answer key text file for all chapter quizzes
    def write_answer_key(path):
        with open(path, "w", encoding="utf-8") as f:
            for num, title, paras, fig_label in chapters:
                f.write(f"Chapter {num}: {title}\n")
                qs = chapter_quiz_questions(num)
                # Provide concise model answers for ch1 and ch2, generic otherwise
                for i, q in enumerate(qs, start=1):
                    ans = ""
                    if num == 1:
                        answers_ch1 = [
                            "A distributed version control system for tracking changes.",
                            "Git is the VCS; GitHub is a hosting/collaboration platform built on Git.",
                            "A storage for project files and their history.",
                            "Commit stores a snapshot of changes, message, author, timestamp.",
                            "A parallel line of development to isolate work.",
                            "Creates a local copy of the remote repository.",
                            "History, collaboration, backup; example: safe experimentation and rollback.",
                            "To explain what changed and why; helps reviewers.",
                            "Use git log or git log --oneline to view history.",
                            "When developing a new feature, to avoid affecting main branch.",
                        ]
                        ans = answers_ch1[i - 1]
                    elif num == 2:
                        answers_ch2 = [
                            "Download from git-scm.com, use package managers (apt, brew, choco), or GitHub Desktop.",
                            'git config --global user.name "Your Name"; git config --global user.email "you@example.com"',
                            "So commits link to your identity and show on GitHub profile.",
                            "A GUI client to manage repos locally and sync with GitHub.",
                            "CLI enables scripting and faster workflows for advanced users.",
                            "Add a remote (HTTPS/SSH) and push; or use GitHub Desktop to connect.",
                            "SSH is key-based and convenient; HTTPS may prompt for credentials or token.",
                            "Run git --version or git --help to verify installation.",
                            "In the global git config file (~/.gitconfig or %USERPROFILE%\\.gitconfig).",
                            "Sign up at github.com and verify your email.",
                        ]
                        ans = answers_ch2[i - 1]
                    else:
                        ans = "See chapter content for the model answer."
                    f.write(f"{i}) Q: {qs[i - 1]}\n   A: {ans}\n")
                f.write("\n")

    answers_path = f"GitHub_Fundamentals_answers_{ts}.txt"
    write_answer_key(answers_path)
    print(f"Created answer key: {answers_path}")


if __name__ == "__main__":
    main()
