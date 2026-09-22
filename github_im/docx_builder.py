from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from github_im.model import Activity, Module, Question
from github_im.ooxml import (
    add_bookmark, add_page_number, add_toc, enable_field_updates,
    prevent_row_split, repeat_table_header, restart_numbered_paragraphs,
    set_cell_margins, set_image_description,
    set_table_borders, shade_cell,
)
from github_im.visuals import COVER_VISUAL, VISUALS


NAVY, PALE_BLUE, PALE_GRAY = "17365D", "EAF2F8", "F2F2F2"
BLACK = RGBColor(0, 0, 0)


def _font(style, name: str, size: float, bold: bool = False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def _configure_styles(document: Document):
    styles = document.styles
    normal = styles["Normal"]
    _font(normal, "Arial", 11)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in (
        ("Title", 20, 0, 12), ("Subtitle", 13, 0, 10),
        ("Heading 1", 16, 14, 8), ("Heading 2", 13, 12, 6),
        ("Heading 3", 11, 10, 4),
    ):
        style = styles[name]
        _font(style, "Arial", size, name != "Subtitle")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    title_border = styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr"))
    if title_border is not None:
        styles["Title"].element.get_or_add_pPr().remove(title_border)
    command = styles["Command"] if "Command" in styles else styles.add_style("Command", WD_STYLE_TYPE.PARAGRAPH)
    _font(command, "Consolas", 9.5)
    command.paragraph_format.left_indent = Inches(0.25)
    command.paragraph_format.right_indent = Inches(0.25)
    command.paragraph_format.space_before = Pt(3)
    command.paragraph_format.space_after = Pt(3)
    command.paragraph_format.keep_together = True
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), PALE_GRAY)
    command.element.get_or_add_pPr().append(shading)
    caption = styles["Caption"]
    _font(caption, "Arial", 9)
    caption.font.italic = True
    caption.paragraph_format.keep_together = True


def _configure_section(section):
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.right_margin = Inches(1)
    section.left_margin = Inches(1.5)
    section.header_distance = section.footer_distance = Inches(0.4)


def _header_footer(section, title: str):
    header = section.header.paragraphs[0]
    header.text = title
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.name, run.font.size = "Arial", Pt(8)
        run.font.color.rgb = RGBColor(89, 89, 89)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)


def _heading(document, text: str, level: int):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _list(document, items, style="List Bullet"):
    paragraphs = [document.add_paragraph(str(item), style=style) for item in items]
    if style == "List Number":
        restart_numbered_paragraphs(document, paragraphs, style)
    return paragraphs


def _format_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    if header and table.rows:
        repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0 and header:
                shade_cell(cell, NAVY)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb, run.bold = RGBColor(255, 255, 255), True
            elif row_index % 2 == 0:
                shade_cell(cell, PALE_BLUE)
    return table


def _questions(document, questions: list[Question]):
    paragraphs = []
    for question in questions:
        paragraphs.append(document.add_paragraph(question.prompt, style="List Number"))
        for choice in question.choices:
            paragraph = document.add_paragraph(choice, style="List Bullet 2")
            paragraph.paragraph_format.left_indent = Inches(0.65)
    restart_numbered_paragraphs(document, paragraphs)


def _activity(document, activity: Activity, heading_level=3):
    _heading(document, activity.title, heading_level)
    _list(document, activity.instructions, "List Number")
    paragraph = document.add_paragraph()
    paragraph.add_run("Required evidence: ").bold = True
    paragraph.add_run(activity.evidence)
    if activity.safety_note:
        paragraph = document.add_paragraph()
        paragraph.add_run("Safety and privacy: ").bold = True
        paragraph.add_run(activity.safety_note)


def _front_matter(document: Document, module: Module, asset_dir: Path):
    overline = document.add_paragraph()
    overline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    overline.paragraph_format.space_after = Pt(8)
    run = overline.add_run("CMU INSTRUCTIONAL MODULE")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(31, 111, 161)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run(module.title)
    run.font.size = Pt(27)
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("A CMU Format Instructional Module")
    artwork = document.add_paragraph()
    artwork.alignment = WD_ALIGN_PARAGRAPH.CENTER
    artwork.paragraph_format.space_before = Pt(8)
    artwork.paragraph_format.space_after = Pt(10)
    shape = artwork.add_run().add_picture(
        str(asset_dir / COVER_VISUAL.filename), height=Inches(5.05)
    )
    set_image_description(shape, COVER_VISUAL.alt_text)
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_before = Pt(3)
    author.add_run("by ").italic = True
    author_run = author.add_run("RINANTE M. BUNTOD")
    author_run.bold = True
    author_run.font.size = Pt(13)
    audience = document.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    audience.paragraph_format.space_before = Pt(8)
    audience.add_run("Semester Edition for ")
    audience.add_run(module.audience).bold = True
    document.add_page_break()
    _heading(document, "Preface", 1)
    for paragraph in module.preface:
        document.add_paragraph(paragraph)
    _heading(document, "Table of Contents", 1)
    add_toc(document.add_paragraph())
    _heading(document, "List of Figures", 1)
    for unit in module.units:
        document.add_paragraph(f"Figure {unit.number}. {VISUALS[unit.figure_key].caption}")
    _heading(document, "Glossary", 1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Term", "Definition"
    for term, definition in sorted(module.glossary.items()):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = term, definition
    _format_table(table)
    _heading(document, "Semester Roadmap", 1)
    table = document.add_table(rows=1, cols=5)
    for cell, label in zip(table.rows[0].cells, ("Unit", "Topic", "Time", "Output", "Assessment")):
        cell.text = label
    for item in module.semester_roadmap:
        cells = table.add_row().cells
        for cell, key in zip(cells, ("unit", "title", "time", "output", "assessment")):
            cell.text = item[key]
    _format_table(table)


def _unit(document: Document, unit, asset_dir: Path):
    heading = _heading(document, f"Unit {unit.number}: {unit.title}", 1)
    heading.paragraph_format.page_break_before = True
    add_bookmark(heading, f"unit_{unit.number}", unit.number)
    paragraph = document.add_paragraph()
    paragraph.add_run("Time allotment: ").bold = True
    paragraph.add_run(unit.time_allotment)
    for label, body in (("Overview", unit.overview), ("Introduction and Rationale", unit.rationale)):
        _heading(document, label, 2)
        document.add_paragraph(body)
    _heading(document, "Learning Objectives", 2)
    _list(document, unit.objectives, "List Number")
    _heading(document, "Learner Instructions", 2)
    _list(document, unit.learner_instructions)
    _heading(document, "Materials", 2)
    _list(document, unit.materials)
    _heading(document, "Pre-Test", 2)
    document.add_paragraph("Answer each item before studying the unit. Check your work only after completing the unit.")
    _questions(document, unit.pretest)
    spec = VISUALS[unit.figure_key]
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(asset_dir / spec.filename), width=Inches(6.3))
    set_image_description(shape, spec.alt_text)
    caption = document.add_paragraph(f"Figure {unit.number}. {spec.caption}", style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    source = document.add_paragraph("Source: Original instructional illustration generated for this module.")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for label, content in (("Concepts and Discussion", unit.discussion),):
        _heading(document, label, 2)
        for item in content:
            document.add_paragraph(item)
    _heading(document, "Command Examples", 2)
    for command, explanation in unit.command_examples:
        document.add_paragraph(command, style="Command")
        document.add_paragraph(explanation)
    _heading(document, "Guided Learning Activity", 2)
    for activity in unit.activities:
        _activity(document, activity)
    _heading(document, "Self-Test", 2)
    _questions(document, unit.self_test)
    _heading(document, "Performance Task", 2)
    _activity(document, unit.performance_task)
    _heading(document, "Post-Test", 2)
    _questions(document, unit.posttest)
    _heading(document, "Suggested Readings", 2)
    _list(document, unit.supplementary_readings)
    _heading(document, "References", 2)
    for reference in unit.references:
        document.add_paragraph(reference)


def _back_matter(document: Document, module: Module):
    heading = _heading(document, "Capstone Project", 1)
    heading.paragraph_format.page_break_before = True
    _activity(document, module.capstone, 2)
    _heading(document, "Assessment Rubrics", 1)
    names = {"unit_performance": "Unit Performance Task Rubric", "pull_request": "Pull Request Rubric", "capstone": "Capstone Rubric"}
    for key, criteria in module.rubrics.items():
        _heading(document, names.get(key, key.replace("_", " ").title()), 2)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Criterion", "Performance descriptors"
        for criterion, descriptor in criteria:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = criterion, descriptor
        _format_table(table)
    heading = _heading(document, "Answer Key", 1)
    heading.paragraph_format.page_break_before = True
    document.add_paragraph("Instructor copy. Accept equivalent wording when it demonstrates the same accurate concept.")
    for unit in module.units:
        _heading(document, f"Unit {unit.number}: {unit.title}", 2)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Item", "Expected answer"
        for item_id, answer in unit.answer_key.items():
            cells = table.add_row().cells
            cells[0].text, cells[1].text = item_id, answer
        _format_table(table)
    _heading(document, "Git Command Cheat Sheet", 1)
    commands = [
        ("git --version", "Verify the Git installation."), ("git config --global user.name \"Name\"", "Set the default commit author name."),
        ("git config --global user.email \"email\"", "Set the default commit author email."), ("git init", "Create a repository in the current folder."),
        ("git clone <url>", "Create a local clone of a remote repository."), ("git status", "Inspect branch, staging, and working-tree state."),
        ("git diff", "Review unstaged changes."), ("git diff --staged", "Review staged changes."), ("git add <path>", "Stage a selected path."),
        ("git commit -m \"message\"", "Record staged content."), ("git log --oneline --graph", "Inspect compact repository history."),
        ("git branch", "List local branches."), ("git switch -c <branch>", "Create and switch to a feature branch."),
        ("git switch <branch>", "Switch to an existing branch."), ("git merge <branch>", "Integrate a branch into the current branch."),
        ("git remote -v", "Inspect configured remotes."), ("git fetch", "Download remote history without integrating it."),
        ("git pull", "Fetch and integrate remote changes."), ("git push -u origin <branch>", "Publish a branch and set its upstream."),
        ("git restore <path>", "Restore an unstaged working-tree file."),
    ]
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Command", "Purpose"
    for command, purpose in commands:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = command, purpose
    _format_table(table)
    _heading(document, "References", 1)
    references = set(module.references)
    for unit in module.units:
        references.update(unit.references)
    for reference in sorted(references):
        document.add_paragraph(reference)
    heading = _heading(document, "About the Author", 1)
    heading.paragraph_format.page_break_before = True
    document.add_paragraph("Rinante M. Buntod is the author of this instructional module. Additional professional and academic biographical information may be added before publication.")


def build_docx(module: Module, asset_dir: Path, output_path: Path) -> Path:
    asset_dir, output_path = Path(asset_dir), Path(output_path)
    document = Document()
    document.core_properties.title = module.title
    document.core_properties.author = "Rinante M. Buntod"
    for section in document.sections:
        _configure_section(section)
        _header_footer(section, module.title)
        section.different_first_page_header_footer = True
        section.first_page_header.paragraphs[0].text = ""
        section.first_page_footer.paragraphs[0].text = ""
    _configure_styles(document)
    enable_field_updates(document)
    _front_matter(document, module, asset_dir)
    for unit in module.units:
        _unit(document, unit, asset_dir)
    _back_matter(document, module)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
