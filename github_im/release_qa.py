from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Mm, Pt
from pypdf import PdfReader

from github_im.content.catalog import build_module


REQUIRED_SECTIONS = (
    "Preface", "Table of Contents", "List of Figures", "Glossary",
    "Semester Roadmap", "Capstone Project", "Assessment Rubrics",
    "Answer Key", "Git Command Cheat Sheet", "References", "About the Author",
)
EXPECTED_FIGURES = 14
FORBIDDEN_PATTERNS = (
    re.compile(r"\bTODO\b", re.IGNORECASE),
    re.compile(r"\bTBD\b", re.IGNORECASE),
    re.compile(r"lorem ipsum", re.IGNORECASE),
    re.compile(r"BEGIN (?:RSA |OPENSSH |EC )?PRIVATE KEY", re.IGNORECASE),
    re.compile(r"ghp_[A-Za-z0-9]{20,}"),
    re.compile(r"github_pat_[A-Za-z0-9_]{20,}"),
)


def _docx_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def _forbidden(text: str) -> bool:
    return any(pattern.search(text) for pattern in FORBIDDEN_PATTERNS)


def audit_docx(path: Path) -> list[str]:
    path = Path(path)
    if not path.exists():
        return ["DOCX file is missing."]
    if path.stat().st_size == 0:
        return ["DOCX file is empty."]
    try:
        document = Document(path)
    except Exception as exc:
        return [f"DOCX cannot be opened: {exc}"]
    errors: list[str] = []
    section = document.sections[0]
    if abs(section.page_width - Mm(210)) > 1000 or abs(section.page_height - Mm(297)) > 1000:
        errors.append("DOCX page size is not A4.")
    expected_margins = (Inches(1), Inches(1), Inches(1.5), Inches(1))
    actual_margins = (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin)
    if actual_margins != expected_margins:
        errors.append("DOCX margins do not match the CMU format.")
    normal = document.styles["Normal"]
    if normal.font.name != "Arial" or normal.font.size != Pt(11) or normal.paragraph_format.line_spacing != 1:
        errors.append("DOCX Normal style does not use Arial 11 single spacing.")
    text = _docx_text(document)
    for section_name in REQUIRED_SECTIONS:
        if section_name not in text:
            errors.append(f"DOCX is missing required section: {section_name}.")
    for number in range(1, 14):
        if f"Unit {number}:" not in text:
            errors.append(f"DOCX is missing Unit {number} heading.")
        if f"Figure {number}." not in text:
            errors.append(f"DOCX is missing Figure {number} caption.")
    if _forbidden(text):
        errors.append("DOCX contains a draft marker or forbidden secret material.")
    if len(document.inline_shapes) != EXPECTED_FIGURES:
        errors.append(
            f"DOCX contains {len(document.inline_shapes)} figures; expected {EXPECTED_FIGURES}."
        )
    try:
        with zipfile.ZipFile(path) as package:
            xml = package.read("word/document.xml").decode("utf-8")
    except Exception as exc:
        errors.append(f"DOCX package cannot be inspected: {exc}")
    else:
        descriptions = re.findall(r'\bdescr="([^"]*)"', xml)
        if (
            len(descriptions) != EXPECTED_FIGURES
            or any(not value.strip() for value in descriptions)
        ):
            errors.append("DOCX figure descriptions are missing or incomplete.")
    module = build_module()
    expected_ids = {
        question.id
        for unit in module.units
        for question in unit.pretest + unit.self_test + unit.posttest
    }
    missing_answers = sorted(item_id for item_id in expected_ids if item_id not in text)
    if missing_answers:
        errors.append(f"DOCX answer key is incomplete; {len(missing_answers)} item identifiers are missing.")
    return errors


def audit_pdf_text(text: str, expected_units: int = 13) -> list[str]:
    errors: list[str] = []
    for number in range(1, expected_units + 1):
        if f"Unit {number}:" not in text:
            errors.append(f"PDF is missing Unit {number} heading.")
        if f"Figure {number}." not in text:
            errors.append(f"PDF is missing Figure {number} caption.")
    for section_name in REQUIRED_SECTIONS:
        if section_name not in text:
            errors.append(f"PDF is missing required section: {section_name}.")
    if _forbidden(text):
        errors.append("PDF contains forbidden secret material.")
    return errors


def audit_pdf(path: Path, expected_units: int = 13) -> list[str]:
    path = Path(path)
    if not path.exists():
        return ["PDF file is missing."]
    if path.stat().st_size == 0:
        return ["PDF file is empty."]
    try:
        reader = PdfReader(path)
    except Exception as exc:
        return [f"PDF cannot be opened: {exc}"]
    if len(reader.pages) == 0:
        return ["PDF contains zero pages."]
    for index, page in enumerate(reader.pages, 1):
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        if abs(width - 595.28) > 2 or abs(height - 841.89) > 2:
            return [f"PDF page {index} is not A4."]
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return audit_pdf_text(text, expected_units)
